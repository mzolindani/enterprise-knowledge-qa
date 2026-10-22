
import pymupdf
import pytest
from docx import Document

from app.ingestion import ingest_file


def create_pdf(path):
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "Workplace Safety Policy")
        page.insert_text((72, 100), "Report all incidents to the safety officer.")

        page = pdf.new_page()
        page.insert_text((72, 72), "Incident Escalation Procedure")

        pdf.save(path)


def create_docx(path):
    document = Document()
    document.add_heading("Workplace Safety Policy", level=1)
    document.add_paragraph(
        "Report all incidents to the safety officer."
    )

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Incident Type"
    table.cell(0, 1).text = "Reporting Contact"

    row = table.add_row()
    row.cells[0].text = "Workplace injury"
    row.cells[1].text = "Safety officer"

    document.save(path)


def test_pdf_ingestion_preserves_page_numbers(tmp_path):
    path = tmp_path / "safety.pdf"
    create_pdf(path)

    document = ingest_file(path)

    assert document.file_type == "pdf"
    assert len(document.pages) == 2
    assert document.pages[0].page_number == 1
    assert document.pages[1].page_number == 2
    assert "Workplace Safety Policy" in document.full_text
    assert len(document.content_hash) == 64


def test_docx_ingestion_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "safety.docx"
    create_docx(path)

    document = ingest_file(path)

    assert document.file_type == "docx"
    assert len(document.pages) == 1
    assert document.pages[0].page_number is None
    assert "Workplace Safety Policy" in document.full_text
    assert "Incident Type" in document.full_text
    assert "Safety officer" in document.full_text
    assert "[TABLE]" in document.full_text


def test_identical_files_have_identical_ids(tmp_path):
    path = tmp_path / "safety.pdf"
    create_pdf(path)

    first = ingest_file(path)
    second = ingest_file(path)

    assert first.document_id == second.document_id
    assert first.content_hash == second.content_hash


def test_missing_file_raises_error(tmp_path):
    with pytest.raises(FileNotFoundError):
        ingest_file(tmp_path / "missing.pdf")


def test_unsupported_file_type_is_rejected(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Some text", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        ingest_file(path)


def test_empty_file_is_rejected(tmp_path):
    path = tmp_path / "empty.pdf"
    path.write_bytes(b"")

    with pytest.raises(ValueError, match="empty"):
        ingest_file(path)


def test_scanned_or_empty_pdf_is_rejected(tmp_path):
    path = tmp_path / "image_only.pdf"

    with pymupdf.open() as pdf:
        pdf.new_page()
        pdf.save(path)

    with pytest.raises(ValueError, match="No extractable text"):
        ingest_file(path)


def test_document_id_stays_stable_when_file_content_changes(tmp_path):
    path = tmp_path / "safety.pdf"
    create_pdf(path)

    first = ingest_file(path)

    # Change the content while keeping the same file path.
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "Updated Workplace Safety Policy")
        page.insert_text((72, 100), "Updated incident reporting procedure.")
        pdf.save(path)

    second = ingest_file(path)

    assert first.document_id == second.document_id
    assert first.content_hash != second.content_hash


def test_identical_content_at_different_paths_has_different_document_ids(
    tmp_path,
):
    first_path = tmp_path / "safety_first.pdf"
    second_path = tmp_path / "safety_second.pdf"

    create_pdf(first_path)

    # Copy the exact bytes to ensure both files are identical.
    second_path.write_bytes(first_path.read_bytes())

    first = ingest_file(first_path)
    second = ingest_file(second_path)

    assert first.document_id != second.document_id
    assert first.content_hash == second.content_hash
