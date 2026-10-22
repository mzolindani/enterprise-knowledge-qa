
import pymupdf  # noqa: I001
import pytest
import tiktoken
from docx import Document
from itertools import pairwise
from app.chunking import chunk_document
from app.ingestion import ingest_file


def create_pdf(path, text):
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_textbox(
            (50, 50, 550, 750),
            text,
            fontsize=10,
        )
        pdf.save(path)


def test_chunking_preserves_document_metadata(tmp_path):
    path = tmp_path / "policy.pdf"
    create_pdf(path, "Workplace safety requires clear reporting procedures.")

    document = ingest_file(path)
    chunks = chunk_document(document, chunk_size=20, overlap=3)

    assert chunks
    assert chunks[0].document_id == document.document_id
    assert chunks[0].document_name == "policy.pdf"
    assert chunks[0].page_number == 1
    assert chunks[0].chunk_index == 0
    assert chunks[0].text
    assert chunks[0].content_hash


def test_long_text_is_split_into_multiple_chunks(tmp_path):
    path = tmp_path / "long.pdf"
    text = " ".join(f"word{i}" for i in range(100))
    create_pdf(path, text)

    document = ingest_file(path)
    chunks = chunk_document(document, chunk_size=20, overlap=5)

    assert len(chunks) > 1
    assert all(chunk.token_count <= 20 for chunk in chunks)


def test_chunk_ids_are_deterministic(tmp_path):
    path = tmp_path / "policy.pdf"
    create_pdf(path, "Employees must follow the documented safety procedures.")

    document = ingest_file(path)

    first = chunk_document(document)
    second = chunk_document(document)

    assert [c.chunk_id for c in first] == [
        c.chunk_id for c in second
    ]


def test_chunking_rejects_invalid_parameters(tmp_path):
    path = tmp_path / "policy.pdf"
    create_pdf(path, "A document with some useful content.")

    document = ingest_file(path)

    with pytest.raises(ValueError):
        chunk_document(document, chunk_size=0)

    with pytest.raises(ValueError):
        chunk_document(document, chunk_size=10, overlap=10)

    with pytest.raises(ValueError):
        chunk_document(document, chunk_size=10, overlap=-1)


def test_chunks_respect_token_limit(tmp_path):
    

    path = tmp_path / "long.pdf"
    text = " ".join(f"word{i}" for i in range(150))
    create_pdf(path, text)

    document = ingest_file(path)
    chunks = chunk_document(document, chunk_size=25, overlap=5)
    encoding = tiktoken.get_encoding("cl100k_base")

    assert len(chunks) > 1

    for chunk in chunks:
        assert len(encoding.encode(chunk.text)) <= 25
        assert chunk.token_count == len(encoding.encode(chunk.text))


def test_long_paragraph_chunks_have_overlapping_context(tmp_path):
    path = tmp_path / "overlap.pdf"
    words = [f"word{i}" for i in range(100)]
    create_pdf(path, " ".join(words))

    document = ingest_file(path)
    chunks = chunk_document(document, chunk_size=20, overlap=5)

    assert len(chunks) > 1

    for previous, current in pairwise(chunks):
        previous_words = previous.text.split()
        current_words = current.text.split()

        common_words = 0
        max_common = min(len(previous_words), len(current_words))

        for count in range(1, max_common + 1):
            if previous_words[-count:] == current_words[:count]:
                common_words = count

        assert common_words > 0

        encoding = tiktoken.get_encoding("cl100k_base")
        overlap_text = " ".join(current_words[:common_words])
        assert len(encoding.encode(overlap_text)) <= 5

        # overlap_text = " ".join(current_words[:common_words])
        # assert len(_ENCODING.encode(overlap_text)) <= 5



def test_docx_heading_is_preserved_as_section_metadata(tmp_path):
    path = tmp_path / "headings.docx"

    doc = Document()
    doc.add_heading("Incident Reporting", level=1)
    doc.add_paragraph(
        "Employees must report all workplace incidents promptly."
    )
    doc.save(path)

    document = ingest_file(path)
    chunks = chunk_document(document)

    assert chunks
    assert any(
        chunk.section == "Incident Reporting"
        for chunk in chunks
    )


def test_table_headers_are_repeated_when_table_is_split(tmp_path):
    path = tmp_path / "table.docx"

    doc = Document()
    doc.add_heading("Incident Contacts", level=1)

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Incident Type"
    table.cell(0, 1).text = "Reporting Contact"

    for index in range(15):
        row = table.add_row()
        row.cells[0].text = f"Incident category {index}"
        row.cells[1].text = f"Contact person {index}"

    doc.save(path)

    document = ingest_file(path)
    chunks = chunk_document(document, chunk_size=30, overlap=5)

    table_chunks = [
        chunk for chunk in chunks
        if "Incident Type" in chunk.text
    ]

    assert len(table_chunks) > 1

    for chunk in table_chunks:
        assert "Reporting Contact" in chunk.text
        assert chunk.section == "Incident Contacts"


def test_invalid_token_overlap_is_rejected(tmp_path):
    path = tmp_path / "small.pdf"
    create_pdf(path, "A short document for validation.")

    document = ingest_file(path)

    with pytest.raises(ValueError, match="overlap"):
        chunk_document(document, chunk_size=20, overlap=20)
